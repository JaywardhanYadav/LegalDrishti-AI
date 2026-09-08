import base64
from pathlib import Path
import docx
import pypdf
from openai import OpenAI

from app.core.config import get_settings


class DocumentExtractionError(Exception):
    pass


def extract_text_from_pdf(file_path: str | Path) -> tuple[str, int]:
    path = Path(file_path)
    if not path.is_file():
        raise DocumentExtractionError(f"PDF file not found at: {path}")

    try:
        reader = pypdf.PdfReader(str(path))

        if reader.is_encrypted:
            raise DocumentExtractionError("PDF file is encrypted or password-protected.")

        pages_text: list[str] = []
        total_pages = len(reader.pages)

        for page_idx, page in enumerate(reader.pages, start=1):
            raw_page_text = page.extract_text() or ""
            cleaned_text = raw_page_text.strip()

            if cleaned_text:
                pages_text.append(f"--- [Page {page_idx}] ---\n{cleaned_text}")

        full_text = "\n\n".join(pages_text)
        return full_text, total_pages

    except pypdf.errors.PdfReadError as exc:
        raise DocumentExtractionError(f"Corrupt or unreadable PDF: {exc}") from exc
    except Exception as exc:
        raise DocumentExtractionError(f"Failed to extract PDF text: {exc}") from exc


def extract_text_from_docx(file_path: str | Path) -> tuple[str, int]:
    path = Path(file_path)
    if not path.is_file():
        raise DocumentExtractionError(f"DOCX file not found at: {path}")

    try:
        doc = docx.Document(str(path))
        content_line: list[str] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                content_line.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    content_line.append(" | ".join(row_text))

        full_text = "\n\n".join(content_line)
        return full_text, 1

    except Exception as exc:
        raise DocumentExtractionError(f"Failed to extract DOCX text: {exc}") from exc


def extract_text_from_txt(file_path: str | Path) -> tuple[str, int]:
    path = Path(file_path)
    if not path.is_file():
        raise DocumentExtractionError(f"TXT file not found at: {path}")
    try:
        try:
            content = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            content = path.read_text(encoding="latin-1")
        return content.strip(), 1
    except Exception as exc:
        raise DocumentExtractionError(f"Failed to read TXT file: {exc}") from exc


def extract_text_from_image(file_path: str | Path) -> tuple[str, int]:
    path = Path(file_path)
    if not path.is_file():
        raise DocumentExtractionError(f"Image file not found at: {path}")

    settings = get_settings()
    if not settings.openai_api_key or not settings.openai_api_key.strip():
        raise DocumentExtractionError(
            "OPENAI_API_KEY is not configured in .env. Vision OCR requires an active API key."
        )

    suffix = path.suffix.lower()
    mime_type = "image/png" if suffix == ".png" else "image/jpeg"

    image_bytes = path.read_bytes()
    base64_encoded = base64.b64encode(image_bytes).decode("utf-8")

    try:
        client = OpenAI(api_key=settings.openai_api_key)
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "You are an expert legal document OCR engine. Accurately transcribe "
                        "all readable text, headers, stamps, handwriting, and tables verbatim "
                        "from this legal document image. Maintain layout and structure. "
                        "Do not summarize, extrapolate, or invent missing words. "
                        "If a section is completely unreadable, mark it as [illegible]."
                    ),
                },
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": "Transcribe this legal document image verbatim:"},
                        {
                            "type": "image_url",
                            "image_url": {
                                "url": f"data:{mime_type};base64,{base64_encoded}",
                            },
                        },
                    ],
                },
            ],
            max_tokens=2000,
        )
        extracted_text = response.choices[0].message.content or ""
        return extracted_text.strip(), 1
    except Exception as exc:
        raise DocumentExtractionError(f"Vision OCR transcription failed: {exc}") from exc


def extract_document_text(file_path: str | Path) -> tuple[str, int]:
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(path)
    elif suffix == ".docx":
        return extract_text_from_docx(path)
    elif suffix == ".txt":
        return extract_text_from_txt(path)
    elif suffix in {".png", ".jpg", ".jpeg"}:
        return extract_text_from_image(path)
    else:
        raise DocumentExtractionError(f"Unsupported file format for extraction: '{suffix}'")
